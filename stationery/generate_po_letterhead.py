#!/usr/bin/env python3
"""Generate the NIYAKRISH INDUSTRIES purchase-order letterhead (.docx).

Company details are taken from the GST registration certificate
(GSTIN 29AAKCN0823D1ZM); colors and logo come from the NIYA brand
identity in this repository.
"""

import os

from docx import Document
from docx.enum.section import WD_HEADER_FOOTER
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

BLUE = "16418C"      # Engineering Blue
ORANGE = "FF7A1A"    # Safety Orange
STEEL = "5B6B7C"     # Steel
INK = "101418"       # Ink
LIGHT = "EEF2F8"     # light blue tint for label cells
GRID = "B9C4D6"      # table grid lines

FONT = "Calibri"

COMPANY = "NIYAKRISH INDUSTRIES PRIVATE LIMITED"
ADDRESS_1 = "Survey No. 428, Hasige Hobli, Huliyur Durga,"
ADDRESS_2 = "Kampalapura, Tumakuru, Karnataka – 572123, India"
GSTIN = "29AAKCN0823D1ZM"

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO = os.path.join(ROOT, "logos", "png", "niya-logo-horizontal.png")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT_PO = os.path.join(HERE, "NIYAKRISH-Purchase-Order-Letterhead.docx")
OUT_LETTER = os.path.join(HERE, "NIYAKRISH-Letterhead.docx")


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def add_run(p, text, size=10, bold=False, color=INK, italic=False):
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return run


def tight(p, before=0, after=0, line=None):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if line is not None:
        fmt.line_spacing = line
    return p


def insert_ordered(parent, element, successors):
    """Insert element before the first successor tag present, else append.

    OOXML property containers (pPr, rPr, tblPr, tcPr) require children in a
    fixed schema order; a plain append puts them out of order and strict
    readers refuse the file.
    """
    for tag in successors:
        nxt = parent.find(qn(tag))
        if nxt is not None:
            nxt.addprevious(element)
            return
    parent.append(element)


PBDR_SUCCESSORS = ("w:shd", "w:tabs", "w:spacing", "w:ind",
                   "w:contextualSpacing", "w:jc", "w:rPr", "w:sectPr")
TBL_BORDERS_SUCCESSORS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
TC_SHD_SUCCESSORS = ("w:noWrap", "w:tcMar", "w:textDirection",
                     "w:tcFitText", "w:vAlign", "w:hideMark")


def p_bottom_border(p, color, size, space=1):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    insert_ordered(pPr, pBdr, PBDR_SUCCESSORS)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    insert_ordered(tcPr, shd, TC_SHD_SUCCESSORS)


def table_borders(table, color=GRID, size=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    insert_ordered(tblPr, borders, TBL_BORDERS_SUCCESSORS)


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    insert_ordered(tblPr, borders, TBL_BORDERS_SUCCESSORS)


def fixed_layout(table, widths):
    table.autofit = False  # emits <w:tblLayout w:type="fixed"/>
    # cell tcW alone is not enough: readers size columns from tblGrid,
    # so the gridCol entries must carry the real widths (in twips)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.emu / 635)))
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


def cell_text(cell, text, size=10, bold=False, color=INK,
              align=WD_ALIGN_PARAGRAPH.LEFT, before=1, after=1):
    p = cell.paragraphs[0]
    p.alignment = align
    tight(p, before=before, after=after)
    add_run(p, text, size=size, bold=bold, color=color)
    return p


def _new_letterhead_doc():
    """A4 document with the NIYAKRISH letterhead header and footer."""
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10)
    style.font.color.rgb = rgb(INK)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    section.top_margin = Mm(34)
    section.bottom_margin = Mm(22)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

    content_width = Mm(210 - 32)

    # ---------------- letterhead header (repeats on every page) ----------
    header = section.header
    htable = header.add_table(rows=1, cols=2, width=content_width)
    no_borders(htable)
    fixed_layout(htable, [Mm(70), Mm(108)])

    logo_cell = htable.cell(0, 0)
    logo_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    lp = logo_cell.paragraphs[0]
    tight(lp)
    lp.add_run().add_picture(LOGO, width=Mm(58))

    info_cell = htable.cell(0, 1)
    info_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    ip = info_cell.paragraphs[0]
    ip.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tight(ip, after=1)
    add_run(ip, COMPANY, size=11, bold=True, color=BLUE)
    for line in (ADDRESS_1, ADDRESS_2):
        p = info_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        tight(p)
        add_run(p, line, size=8.5, color=STEEL)
    p = info_cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tight(p)
    add_run(p, "GSTIN: ", size=8.5, bold=True, color=STEEL)
    add_run(p, GSTIN, size=8.5, color=STEEL)
    add_run(p, "   |   CIN: [CIN Number]", size=8.5, color=STEEL)
    p = info_cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tight(p)
    add_run(p, "Phone: [+91-XXXXX XXXXX]   |   Email: [email address]",
            size=8.5, color=STEEL)

    rule1 = header.add_paragraph()
    tight(rule1)
    rule1.paragraph_format.line_spacing = Pt(2)
    add_run(rule1, "", size=1)
    p_bottom_border(rule1, BLUE, 18)
    rule2 = header.add_paragraph()
    tight(rule2)
    rule2.paragraph_format.line_spacing = Pt(2)
    add_run(rule2, "", size=1)
    p_bottom_border(rule2, ORANGE, 6)

    # ---------------- footer -------------------------------------------
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(fp, before=2)
    p_bottom_border(fp, ORANGE, 6, space=0)
    add_run(fp,
            f"Regd. Office: {ADDRESS_1} {ADDRESS_2}",
            size=7.5, color=STEEL)
    fp2 = footer.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(fp2)
    add_run(fp2, f"GSTIN: {GSTIN}", size=7.5, color=STEEL)
    add_run(fp2, "   |   STRENGTH DELIVERED", size=7.5, bold=True,
            color=ORANGE)

    # The orange rule must sit above the footer text: move the border to
    # the top edge instead of the bottom.
    pPr = fp._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    bottom = pBdr.find(qn("w:bottom"))
    bottom.tag = qn("w:top")

    return doc


def build_po():
    doc = _new_letterhead_doc()

    # ---------------- title ---------------------------------------------
    title = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(title, before=2, after=10)
    run = add_run(title, "PURCHASE ORDER", size=16, bold=True, color=BLUE)
    rPr = run._r.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "60")
    insert_ordered(rPr, spacing, ("w:w", "w:kern", "w:position", "w:sz",
                                  "w:szCs", "w:u", "w:vertAlign", "w:lang"))

    # ---------------- PO meta -------------------------------------------
    meta = doc.add_table(rows=2, cols=4)
    table_borders(meta)
    fixed_layout(meta, [Mm(32), Mm(57), Mm(32), Mm(57)])
    meta_fields = [
        ("PO Number", "NIPL/PO/25-26/______", "PO Date", "____________"),
        ("Quotation Ref.", "____________", "Delivery Date", "____________"),
    ]
    for r, (l1, v1, l2, v2) in enumerate(meta_fields):
        cell_text(meta.cell(r, 0), l1, size=9.5, bold=True, color=BLUE)
        shade(meta.cell(r, 0), LIGHT)
        cell_text(meta.cell(r, 1), v1, size=9.5)
        cell_text(meta.cell(r, 2), l2, size=9.5, bold=True, color=BLUE)
        shade(meta.cell(r, 2), LIGHT)
        cell_text(meta.cell(r, 3), v2, size=9.5)

    tight(doc.add_paragraph(), after=4)

    # ---------------- vendor / ship-to ----------------------------------
    vt = doc.add_table(rows=2, cols=2)
    table_borders(vt)
    fixed_layout(vt, [Mm(89), Mm(89)])

    for c, heading in enumerate(("VENDOR / SUPPLIER",
                                 "SHIP TO / DELIVERY ADDRESS")):
        cell = vt.cell(0, c)
        shade(cell, BLUE)
        cell_text(cell, heading, size=9.5, bold=True, color="FFFFFF",
                  before=2, after=2)

    vendor = vt.cell(1, 0)
    vendor_lines = [
        ("M/s. ", "______________________________"),
        ("Address: ", "______________________________"),
        ("", "______________________________"),
        ("GSTIN: ", "______________________________"),
        ("Contact Person: ", "______________________"),
        ("Phone: ", "______________________________"),
    ]
    first = True
    for label, blank in vendor_lines:
        p = vendor.paragraphs[0] if first else vendor.add_paragraph()
        first = False
        tight(p, before=2, after=2)
        if label:
            add_run(p, label, size=9.5, bold=True, color=STEEL)
        add_run(p, blank, size=9.5)

    shipto = vt.cell(1, 1)
    ship_lines = [
        (COMPANY, True, BLUE),
        (ADDRESS_1, False, INK),
        (ADDRESS_2, False, INK),
        (f"GSTIN: {GSTIN}", False, INK),
        ("Contact Person: ______________________", False, INK),
        ("Phone: ______________________________", False, INK),
    ]
    first = True
    for text, bold, color in ship_lines:
        p = shipto.paragraphs[0] if first else shipto.add_paragraph()
        first = False
        tight(p, before=2, after=2)
        add_run(p, text, size=9.5, bold=bold, color=color)

    tight(doc.add_paragraph(), after=4)

    # ---------------- items table ---------------------------------------
    widths = [Mm(11), Mm(63), Mm(18), Mm(15), Mm(15), Mm(26), Mm(30)]
    items = doc.add_table(rows=8, cols=7)
    table_borders(items)
    fixed_layout(items, widths)

    headers = ["Sl. No.", "Description of Goods / Services", "HSN / SAC",
               "Qty", "UOM", "Rate (₹)", "Amount (₹)"]
    for c, h in enumerate(headers):
        cell = items.cell(0, c)
        shade(cell, BLUE)
        cell_text(cell, h, size=9, bold=True, color="FFFFFF",
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=2)

    # repeat the column headings if the table spills onto a new page
    trPr = items.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))

    for r in range(1, 8):
        cell_text(items.cell(r, 0), str(r), size=9,
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=4)
        for c in range(1, 7):
            cell_text(items.cell(r, c), "", size=9, before=4, after=4)

    # totals rows appended to the same grid
    totals = [
        ("Sub Total", False),
        ("CGST @ ____ %", False),
        ("SGST @ ____ %", False),
        ("IGST @ ____ %", False),
        ("Freight / Other Charges", False),
        ("GRAND TOTAL", True),
    ]
    for label, is_grand in totals:
        row = items.add_row()
        for i, w in enumerate(widths):
            row.cells[i].width = w
        label_cell = row.cells[0].merge(row.cells[5])
        amount_cell = row.cells[6]
        cell_text(label_cell, label, size=9, bold=True,
                  color=BLUE if is_grand else INK,
                  align=WD_ALIGN_PARAGRAPH.RIGHT, before=2, after=2)
        cell_text(amount_cell, "", size=9, before=2, after=2)
        if is_grand:
            shade(label_cell, LIGHT)
            shade(amount_cell, LIGHT)

    words = items.add_row()
    for i, w in enumerate(widths):
        words.cells[i].width = w
    wc = words.cells[0].merge(words.cells[6])
    p = wc.paragraphs[0]
    tight(p, before=2, after=2)
    add_run(p, "Amount in Words (₹): ", size=9, bold=True, color=BLUE)
    add_run(p, "_" * 68, size=9)

    tight(doc.add_paragraph(), after=2)

    # ---------------- terms & conditions ---------------------------------
    th = doc.add_paragraph()
    tight(th, before=4, after=3)
    th.paragraph_format.keep_with_next = True
    add_run(th, "TERMS & CONDITIONS", size=10.5, bold=True, color=BLUE)
    p_bottom_border(th, GRID, 4)

    terms = [
        "Please quote the Purchase Order number on all invoices, delivery "
        "challans, packing lists and correspondence.",
        "Goods must be delivered to the Ship To address on or before the "
        "delivery date stated above; time is the essence of this order.",
        "All materials are subject to inspection and approval at the "
        "delivery site. Rejected goods will be returned at the supplier's "
        "risk and cost.",
        "The tax invoice must conform to GST rules and clearly state the "
        "supplier's GSTIN, HSN/SAC codes and this PO number. An e-way bill "
        "must accompany the consignment where applicable.",
        "Payment will be released as per the agreed payment terms from the "
        "date of receipt and acceptance of goods along with a correct tax "
        "invoice.",
        "Prices stated in this Purchase Order are firm. No escalation will "
        "be accepted unless agreed in writing.",
        f"{COMPANY} reserves the right to cancel this order, in whole or "
        "in part, for delayed, short or non-conforming supply.",
        "Any dispute arising out of this order is subject to the exclusive "
        "jurisdiction of the courts at Tumakuru, Karnataka.",
    ]
    for i, t in enumerate(terms, 1):
        p = doc.add_paragraph()
        tight(p, after=2)
        p.paragraph_format.left_indent = Mm(6)
        p.paragraph_format.first_line_indent = Mm(-6)
        add_run(p, f"{i}.  ", size=8.5, bold=True, color=STEEL)
        add_run(p, t, size=8.5, color=INK)

    tight(doc.add_paragraph(), after=8)

    # ---------------- signatures -----------------------------------------
    sig = doc.add_table(rows=1, cols=2)
    no_borders(sig)
    fixed_layout(sig, [Mm(89), Mm(89)])

    left = sig.cell(0, 0)
    p = left.paragraphs[0]
    tight(p, after=42)
    add_run(p, "Prepared / Checked By", size=9.5, bold=True, color=STEEL)
    p = left.add_paragraph()
    tight(p)
    add_run(p, "Name & Signature", size=9, color=STEEL)
    p_bottom_border(p, INK, 4)

    right = sig.cell(0, 1)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tight(p, after=42)
    add_run(p, f"For {COMPANY}", size=9.5, bold=True, color=BLUE)
    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tight(p)
    add_run(p, "Authorised Signatory", size=9, bold=True, color=INK)
    p_bottom_border(p, INK, 4)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(note, before=10)
    add_run(note,
            "This Purchase Order is valid only when issued on the official "
            "letterhead and signed by an authorised signatory.",
            size=7.5, italic=True, color=STEEL)

    doc.save(OUT_PO)
    print(f"Saved: {OUT_PO}")


def build_letter():
    """Blank correspondence letterhead with a standard letter skeleton."""
    doc = _new_letterhead_doc()

    meta = doc.add_table(rows=1, cols=2)
    no_borders(meta)
    fixed_layout(meta, [Mm(89), Mm(89)])
    cell_text(meta.cell(0, 0), "Ref. No.: NIPL/______/25-26", size=10)
    cell_text(meta.cell(0, 1), "Date: ______________", size=10,
              align=WD_ALIGN_PARAGRAPH.RIGHT)

    p = doc.add_paragraph()
    tight(p, before=14, after=2)
    add_run(p, "To,", size=10)
    for _ in range(3):
        p = doc.add_paragraph()
        tight(p, after=2)
        add_run(p, "_______________________________", size=10, color=STEEL)

    p = doc.add_paragraph()
    tight(p, before=10, after=2)
    add_run(p, "Subject: ", size=10, bold=True)
    add_run(p, "_" * 76, size=10, color=STEEL)

    p = doc.add_paragraph()
    tight(p, before=10, after=2)
    add_run(p, "Dear Sir / Madam,", size=10)

    p = doc.add_paragraph()
    tight(p, before=8, after=2)
    add_run(p, "[Type the letter body here]", size=10, italic=True,
            color=STEEL)
    for _ in range(10):
        tight(doc.add_paragraph(), after=2)

    p = doc.add_paragraph()
    tight(p, before=8, after=2)
    add_run(p, "Thanking you,", size=10)
    p = doc.add_paragraph()
    tight(p, before=6, after=2)
    add_run(p, "Yours faithfully,", size=10)
    p = doc.add_paragraph()
    tight(p, after=42)
    add_run(p, f"For {COMPANY}", size=10, bold=True, color=BLUE)
    p = doc.add_paragraph()
    tight(p)
    add_run(p, "Authorised Signatory", size=10, bold=True)
    p = doc.add_paragraph()
    tight(p, before=4)
    add_run(p, "Name: ____________________    "
               "Designation: ____________________", size=9.5, color=STEEL)
    p = doc.add_paragraph()
    tight(p, before=8)
    add_run(p, "Encl.: ", size=9.5, bold=True, color=STEEL)
    add_run(p, "______________________________", size=9.5, color=STEEL)

    doc.save(OUT_LETTER)
    print(f"Saved: {OUT_LETTER}")


if __name__ == "__main__":
    build_po()
    build_letter()
